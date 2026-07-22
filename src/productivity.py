# src/productivity.py
# BSL SL Automation — Productivity Template Population
# Fix: Compatible table access for all openpyxl versions
# Fix: Unique table names across all sheets
# Fix: Re-create tables after paste to restore pivot references [5]

import io
import copy
import re
import pandas as pd
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
import logging

logger = logging.getLogger("BSL_SL")

SHEET_PENETRATION        = "Penetration Per Day"
SHEET_EARLY              = "Stat Result - BL Early SL"
SHEET_REMEDIAL           = "Stat Result - BL Remedial SL"
STAT_RESULT_HEADER_KEYS  = ["Date", "Time", "Name", "LAN", "Status"]
PENETRATION_HEADER_KEYS  = ["DATE", "CLIENT", "ACCOUNTS", "TOTAL DIALED"]


def _find_sheet_exact(wb, exact_names: list) -> str:
    """
    Finds sheet by exact name first, then case-insensitive,
    then partial match.
    """
    logger.info(f"Available sheets: {wb.sheetnames}")
    for name in exact_names:
        if name in wb.sheetnames:
            logger.info(f"Exact match: '{name}'")
            return name
    for name in exact_names:
        for sheet in wb.sheetnames:
            if name.lower() == sheet.lower():
                logger.info(f"Case-insensitive match: '{sheet}'")
                return sheet
    for name in exact_names:
        for sheet in wb.sheetnames:
            if name.lower() in sheet.lower():
                logger.info(f"Partial match '{name}' -> '{sheet}'")
                return sheet
    raise ValueError(
        f"Sheet not found from {exact_names}. "
        f"Available: {wb.sheetnames}"
    )


def _find_header_row_strict(ws, required_keys: list) -> int:
    """
    Finds the FIRST row where ALL required_keys appear as cell values.
    Returns 1-based row number. Returns 1 if not found.
    """
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True),
        start=1
    ):
        row_vals = [
            str(v).strip() if v is not None else ""
            for v in row
        ]
        if all(k in row_vals for k in required_keys):
            logger.info(f"Header row at row {row_idx} in '{ws.title}'")
            return row_idx
    logger.warning(
        f"Header row with {required_keys} not found in '{ws.title}'. "
        f"Defaulting to row 1."
    )
    return 1


def _find_last_data_row_strict(ws, header_row: int) -> int:
    """
    Finds the ACTUAL last row with data below header_row.
    Checks column A for non-empty values.
    """
    last_data_row = header_row
    for row_idx in range(header_row + 1, ws.max_row + 1):
        cell_val = ws.cell(row=row_idx, column=1).value
        if cell_val is not None and str(cell_val).strip() != "":
            last_data_row = row_idx
    logger.info(
        f"Last data row in '{ws.title}': {last_data_row} "
        f"(header at row {header_row})"
    )
    return last_data_row


def _copy_row_format(ws, source_row: int, target_row: int):
    """
    Copies cell formatting from source_row to target_row.
    Fixes ########## display in Penetration Per Day [1].
    """
    for col_idx in range(1, ws.max_column + 1):
        src = ws.cell(row=source_row, column=col_idx)
        tgt = ws.cell(row=target_row, column=col_idx)
        tgt.number_format = src.number_format
        if src.font:
            tgt.font      = copy.copy(src.font)
        if src.fill:
            tgt.fill      = copy.copy(src.fill)
        if src.border:
            tgt.border    = copy.copy(src.border)
        if src.alignment:
            tgt.alignment = copy.copy(src.alignment)
    if ws.row_dimensions[source_row].height:
        ws.row_dimensions[target_row].height = (
            ws.row_dimensions[source_row].height
        )
    logger.info(f"Copied format row {source_row} -> row {target_row}")


def _get_table_info(ws) -> list:
    """
    Gets existing table info (name, style) before removal.
    Compatible with all openpyxl versions.
    """
    table_info = []

    for item in ws.tables.values():
        # Handle string (some openpyxl versions return table name as string)
        if isinstance(item, str):
            info = {
                "name":             item,
                "display_name":     item,
                "style_name":       "TableStyleMedium9",
                "show_first_col":   False,
                "show_last_col":    False,
                "show_row_stripes": True,
                "show_col_stripes": False,
            }
        else:
            # Table object
            info = {
                "name":         getattr(item, "displayName", "DataTable"),
                "display_name": getattr(item, "displayName", "DataTable"),
            }
            style = getattr(item, "tableStyleInfo", None)
            if style:
                info["style_name"]         = getattr(style, "name", "TableStyleMedium9")
                info["show_first_col"]     = getattr(style, "showFirstColumn", False)
                info["show_last_col"]      = getattr(style, "showLastColumn", False)
                info["show_row_stripes"]   = getattr(style, "showRowStripes", True)
                info["show_col_stripes"]   = getattr(style, "showColumnStripes", False)
            else:
                info["style_name"]         = "TableStyleMedium9"
                info["show_first_col"]     = False
                info["show_last_col"]      = False
                info["show_row_stripes"]   = True
                info["show_col_stripes"]   = False

        table_info.append(info)
        logger.info(f"Captured table: {info}")

    return table_info


def _remove_tables(ws) -> list:
    """
    Captures table info then removes ALL tables from worksheet.
    Compatible with all openpyxl versions.
    Returns captured table info for re-creation after pasting.
    """
    if not ws.tables:
        logger.info(f"No tables in '{ws.title}'.")
        return []

    table_info  = _get_table_info(ws)
    table_keys  = list(ws.tables.keys())
    logger.info(
        f"Removing {len(table_keys)} table(s) "
        f"from '{ws.title}': {table_keys}"
    )
    for key in table_keys:
        del ws.tables[key]
    logger.info(f"Tables removed from '{ws.title}'.")
    return table_info


def _recreate_table(
    ws,
    table_info: list,
    header_row: int,
    last_data_row: int,
    num_cols: int
):
    """
    Re-creates Excel table(s) with updated range after pasting.
    Generates UNIQUE table name across ALL sheets in workbook.
    Restores pivot table data source references [5].
    """
    start_col   = get_column_letter(1)
    end_col     = get_column_letter(num_cols)
    table_range = f"{start_col}{header_row}:{end_col}{last_data_row}"

    # Collect ALL existing table names across ALL sheets
    existing_names = set()
    for sheetname in ws.parent.sheetnames:
        for tbl_name in ws.parent[sheetname].tables.keys():
            existing_names.add(tbl_name.lower())
    logger.info(f"Existing table names: {existing_names}")

    if not table_info:
        table_info = [{
            "name":             "DataTable",
            "display_name":     "DataTable",
            "style_name":       "TableStyleMedium9",
            "show_first_col":   False,
            "show_last_col":    False,
            "show_row_stripes": True,
            "show_col_stripes": False,
        }]

    for info in table_info:
        # Sanitize base name
        base_name = re.sub(
            r"[^A-Za-z0-9_]", "_",
            info.get("name", "DataTable")
        )
        if base_name and base_name[0].isdigit():
            base_name = f"T_{base_name}"
        if not base_name:
            base_name = "DataTable"

        # Generate unique name
        unique_name = base_name
        counter     = 1
        while unique_name.lower() in existing_names:
            unique_name = f"{base_name}_{counter}"
            counter    += 1

        existing_names.add(unique_name.lower())

        style = TableStyleInfo(
            name=info.get("style_name", "TableStyleMedium9"),
            showFirstColumn=info.get("show_first_col", False),
            showLastColumn=info.get("show_last_col", False),
            showRowStripes=info.get("show_row_stripes", True),
            showColumnStripes=info.get("show_col_stripes", False),
        )
        tbl                = Table(
            displayName=unique_name,
            ref=table_range
        )
        tbl.tableStyleInfo = style
        ws.add_table(tbl)
        logger.info(
            f"Re-created table '{unique_name}' "
            f"with range '{table_range}' in '{ws.title}'"
        )


def _clear_below_header(ws, header_row: int):
    """
    Clears ALL cell values below header_row only.
    """
    logger.info(
        f"Clearing rows {header_row + 1} to {ws.max_row} "
        f"in '{ws.title}'"
    )
    for row in ws.iter_rows(
        min_row=header_row + 1,
        max_row=ws.max_row
    ):
        for cell in row:
            cell.value = None


def _paste_below_header(ws, df: pd.DataFrame, header_row: int) -> int:
    """
    Pastes DataFrame starting at header_row + 1, column A.
    Pure positional paste (A B C D...).
    Skips completely empty rows.
    Returns the actual last row written.
    """
    start_row  = header_row + 1
    actual_row = start_row
    logger.info(
        f"Pasting {len(df)} rows into '{ws.title}' "
        f"at row {start_row}"
    )
    for row_data in df.itertuples(index=False):
        values = [
            str(v).strip() if v is not None else ""
            for v in row_data
        ]
        if not any(v != "" for v in values):
            continue
        for c_idx, value in enumerate(row_data, start=1):
            ws.cell(row=actual_row, column=c_idx).value = value
        actual_row += 1

    last_written = actual_row - 1
    logger.info(
        f"Rows pasted: {last_written - start_row + 1} "
        f"(rows {start_row} to {last_written})"
    )
    return last_written


def populate_productivity(
    template_file,
    dialer_df: pd.DataFrame,
    early_df: pd.DataFrame,
    remedial_df: pd.DataFrame
) -> io.BytesIO:
    """
    Populates the Productivity Template [4][5]:
    1. Penetration Per Day  — remove table → append → re-create
    2. Stat Result Remedial — remove table → clear → paste → re-create
    3. Stat Result Early SL — remove table → clear → paste → re-create
    """
    logger.info("Loading Productivity Template...")
    wb = openpyxl.load_workbook(template_file)
    logger.info(f"Sheets found: {wb.sheetnames}")

    # ------------------------------------------------------------------ #
    # STEP 1: Penetration Per Day
    # ------------------------------------------------------------------ #
    pen_sheet      = _find_sheet_exact(wb, [
        "Penetration Per Day", "Penetration", "Per Day"
    ])
    ws_pen         = wb[pen_sheet]
    pen_table_info = _remove_tables(ws_pen)

    try:
        pen_header_row = _find_header_row_strict(
            ws_pen, PENETRATION_HEADER_KEYS
        )
    except ValueError:
        try:
            pen_header_row = _find_header_row_strict(
                ws_pen, ["Date", "Client", "Accounts"]
            )
        except ValueError:
            pen_header_row = 1
            logger.warning("Penetration header not found. Using row 1.")

    pen_last_row = _find_last_data_row_strict(ws_pen, pen_header_row)
    pen_next_row = pen_last_row + 1

    if pen_last_row > pen_header_row:
        _copy_row_format(ws_pen, pen_last_row, pen_next_row)

    for c_idx, value in enumerate(dialer_df.iloc[0].values, start=1):
        ws_pen.cell(row=pen_next_row, column=c_idx).value = value

    _recreate_table(
        ws_pen, pen_table_info,
        pen_header_row, pen_next_row,
        len(dialer_df.columns)
    )
    logger.info(f"Dialer appended at row {pen_next_row}.")

    # ------------------------------------------------------------------ #
    # STEP 2: Stat Result - BL Remedial SL
    # ------------------------------------------------------------------ #
    rem_sheet      = _find_sheet_exact(wb, [
        "Stat Result - BL Remedial SL",
        "Stat Result - BL Remedial",
        "BL Remedial SL",
        "Remedial SL",
        "Remedial"
    ])
    ws_rem         = wb[rem_sheet]
    rem_table_info = _remove_tables(ws_rem)

    try:
        rem_header_row = _find_header_row_strict(
            ws_rem, STAT_RESULT_HEADER_KEYS
        )
    except ValueError:
        rem_header_row = 1
        logger.warning("Remedial header not found. Using row 1.")

    _clear_below_header(ws_rem, rem_header_row)
    rem_last = _paste_below_header(ws_rem, remedial_df, rem_header_row)
    _recreate_table(
        ws_rem, rem_table_info,
        rem_header_row,
        max(rem_last, rem_header_row + 1),
        len(remedial_df.columns)
    )
    logger.info(f"Remedial pasted: {len(remedial_df)} rows.")

    # ------------------------------------------------------------------ #
    # STEP 3: Stat Result - BL Early SL
    # ------------------------------------------------------------------ #
    early_sheet      = _find_sheet_exact(wb, [
        "Stat Result - BL Early SL",
        "Stat Result - BL Early",
        "BL Early SL",
        "Early SL",
        "Early"
    ])
    ws_early         = wb[early_sheet]
    early_table_info = _remove_tables(ws_early)

    try:
        early_header_row = _find_header_row_strict(
            ws_early, STAT_RESULT_HEADER_KEYS
        )
    except ValueError:
        early_header_row = 1
        logger.warning("Early SL header not found. Using row 1.")

    _clear_below_header(ws_early, early_header_row)
    early_last = _paste_below_header(ws_early, early_df, early_header_row)
    _recreate_table(
        ws_early, early_table_info,
        early_header_row,
        max(early_last, early_header_row + 1),
        len(early_df.columns)
    )
    logger.info(f"Early SL pasted: {len(early_df)} rows.")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    logger.info("Productivity Template populated successfully.")
    return output